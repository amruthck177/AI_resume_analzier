import io
import re
from typing import List, Tuple
from pypdf import PdfReader
from docx import Document
from .models import ParsedDocument, DocumentSection, DocumentItem


KNOWN_SECTIONS = [
    "EXPERIENCE", "WORK EXPERIENCE", "EMPLOYMENT", "PROFESSIONAL EXPERIENCE",
    "EDUCATION", "ACADEMIC BACKGROUND",
    "SKILLS", "TECHNICAL SKILLS", "CORE COMPETENCIES", "TECHNOLOGIES",
    "PROJECTS", "PERSONAL PROJECTS", "KEY PROJECTS",
    "CERTIFICATIONS", "LICENSES", "AWARDS", "ACHIEVEMENTS",
    "SUMMARY", "PROFESSIONAL SUMMARY", "PROFILE", "ABOUT ME",
    "PUBLICATIONS", "VOLUNTEER"
]


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])


def parse_resume_text(raw_text: str) -> ParsedDocument:
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    if not lines:
        return ParsedDocument(
            candidateName="Resume Candidate",
            contactInfo="No contact info provided",
            sections=[],
            rawText=""
        )

    # First line is typically candidate name
    candidate_name = lines[0].replace("#", "").strip()
    contact_info = ""
    start_idx = 1

    # Look for contact info in the next few lines
    contact_candidates = []
    while start_idx < min(len(lines), 4):
        line = lines[start_idx]
        # Check if line looks like contact info (contains email, phone, pipe, bullet, dot)
        if any(c in line for c in ["@", "+", "linkedin", "github", "·", "|", ","]):
            contact_candidates.append(line)
            start_idx += 1
        elif any(sec in line.upper() for sec in KNOWN_SECTIONS):
            break
        else:
            if not contact_candidates:
                contact_candidates.append(line)
            start_idx += 1

    contact_info = " · ".join(contact_candidates) if contact_candidates else ""

    # Parse sections
    sections: List[DocumentSection] = []
    current_sec_title = "SUMMARY"
    current_items: List[DocumentItem] = []
    sec_counter = 1
    item_counter = 1

    def flush_section():
        nonlocal current_sec_title, current_items, sec_counter
        if current_items:
            sec_id = f"sec-{re.sub(r'[^a-zA-Z0-9]', '', current_sec_title.lower())}-{sec_counter}"
            sections.append(DocumentSection(
                id=sec_id,
                title=current_sec_title,
                items=current_items
            ))
            sec_counter += 1
            current_items = []

    for i in range(start_idx, len(lines)):
        line = lines[i]
        upper_line = line.upper().strip()

        # Check if this line is a section header
        is_header = False
        matched_title = ""
        for known in KNOWN_SECTIONS:
            # Check exact match or short heading like 'EXPERIENCE' or '## Experience'
            clean_test = re.sub(r'^[#*\s\-_:]+|[#*\s\-_:]+$', '', upper_line)
            if clean_test == known or clean_test.startswith(known + ":"):
                is_header = True
                matched_title = known
                break

        if is_header:
            flush_section()
            current_sec_title = matched_title
            item_counter = 1
            continue

        # Item categorization
        item_id = f"item-{re.sub(r'[^a-zA-Z0-9]', '', current_sec_title.lower())}-{item_counter}"
        item_counter += 1

        is_bullet = bool(re.match(r'^[\u2022\u2023\u25E6\u2043\u2219\-\*]\s+', line))
        clean_text = re.sub(r'^[\u2022\u2023\u25E6\u2043\u2219\-\*]\s+', '', line)

        # Check if it looks like a role/subheading (e.g. "Software Engineer, Google (2022-2024)")
        if not is_bullet and (any(d in line for d in ["202", "201", "199", "Present", "–", "-"]) or len(line) < 60 and not line.endswith('.')):
            current_items.append(DocumentItem(
                id=item_id,
                text=line,
                itemType="subheading"
            ))
        elif is_bullet:
            current_items.append(DocumentItem(
                id=item_id,
                text=clean_text,
                itemType="bullet"
            ))
        else:
            current_items.append(DocumentItem(
                id=item_id,
                text=line,
                itemType="text"
            ))

    flush_section()

    return ParsedDocument(
        candidateName=candidate_name,
        contactInfo=contact_info,
        sections=sections,
        rawText=raw_text
    )
